"""Shared workbook validation and display groups for Python and browser plots."""

import math
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
GROUPS = (
    ("Vibrio-related taxa", ("Vibrio", "Photobacterium", "Shewanella"), "#d7191c"),
    ("Bacillus-related taxa", ("Bacillus", "Lactobacillus"), "#1476a3"),
    ("Other bacterial taxa", (), "#72a936"),
)
TISSUE_NAMES = {"HP": "Hepatopancreas (HP)", "Gut": "Gut"}


def read_data(path):
    """Read summary rows; reject ambiguous keys, invalid counts or stale freq."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "summary" not in workbook.sheetnames:
            raise ValueError("The workbook must contain a 'summary' sheet.")
        values = iter(workbook["summary"].values)
        headers = next(values, ())
        required = ("Time", "Tissue", "Taxon", "Treatment", "n_shrimp", "n_pos", "freq")
        if any(headers.count(name) != 1 for name in required):
            raise ValueError(f"Required unique columns: {', '.join(required)}")
        records, seen, denominators = [], set(), {}
        for row_number, values_row in enumerate(values, start=2):
            if all(value is None for value in values_row):
                continue
            row = dict(zip(headers, values_row))
            for name in required[:4]:
                value = row[name]
                if not isinstance(value, str) or not value.strip():
                    raise ValueError(f"Row {row_number}: missing or invalid {name}.")
                row[name] = value.strip()
            row.setdefault("n_missing", 0)
            numeric = ["n_shrimp", "n_missing"]
            numeric += ["n_pos_observed"] if row["n_missing"] else ["n_pos", "freq"]
            if not row["n_missing"] and "n_pos_observed" in row:
                numeric.append("n_pos_observed")
            for name in numeric:
                value = row.get(name)
                if (isinstance(value, bool) or not isinstance(value, (int, float))
                        or not math.isfinite(value)):
                    raise ValueError(f"Row {row_number}: {name} must be a finite number.")
            n, positive, missing = row["n_shrimp"], row["n_pos"], row["n_missing"]
            if n <= 0 or int(n) != n or int(missing) != missing or not 0 <= missing <= n:
                raise ValueError(f"Row {row_number}: invalid sample size or missing count.")
            if missing:
                observed = row["n_pos_observed"]
                if positive is not None or row["freq"] is not None or int(observed) != observed or not 0 <= observed <= n - missing:
                    raise ValueError(f"Row {row_number}: incomplete measurements require unknown n_pos/freq and valid observed positives.")
                row["prevalence_pct"] = None
            elif not 0 <= positive <= n or int(positive) != positive:
                raise ValueError(f"Row {row_number}: require integer 0 <= n_pos <= n_shrimp, n_shrimp > 0.")
            elif not math.isclose(row["freq"], positive / n, rel_tol=0, abs_tol=1e-9):
                raise ValueError(f"Row {row_number}: freq does not equal n_pos / n_shrimp.")
            else:
                if row.get("n_pos_observed", positive) != positive:
                    raise ValueError(f"Row {row_number}: observed positives disagree with n_pos.")
                row["n_pos_observed"] = positive
                row["prevalence_pct"] = 100 * positive / n
            key = tuple(row[name] for name in required[:4])
            if key in seen:
                raise ValueError(f"Row {row_number}: duplicate Time/Tissue/Taxon/Treatment: {key}")
            seen.add(key)
            cohort = (row["Time"], row["Tissue"], row["Treatment"])
            if denominators.setdefault(cohort, n) != n:
                raise ValueError(f"Row {row_number}: inconsistent n_shrimp for {cohort}.")
            records.append(row)
        if not records:
            raise ValueError("The summary sheet has no data rows.")
        return records
    finally:
        workbook.close()


def taxon_group(taxon):
    """Visual bins follow the reference figure, not a formal taxonomic rank."""
    genus = taxon.split()[0]
    return next((i for i, (_, genera, _) in enumerate(GROUPS) if genus in genera), 2)
