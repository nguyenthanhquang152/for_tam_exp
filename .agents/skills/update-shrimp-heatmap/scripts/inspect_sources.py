"""Read-only workbook profiles; run with the project's Python environment."""

import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def profile(path, max_cells):
    path = Path(path)
    workbook = load_workbook(path, data_only=False)
    try:
        sheets = []
        for sheet in workbook:
            complete = sheet.max_row * sheet.max_column <= max_cells
            preview = [[cell.value for cell in row] for row in sheet.iter_rows(
                max_row=min(sheet.max_row, 8), max_col=min(sheet.max_column, 24))]
            item = {"name": sheet.title, "declaredRows": sheet.max_row,
                    "declaredColumns": sheet.max_column,
                    "mergedRanges": [str(area) for area in sheet.merged_cells.ranges],
                    "preview": preview, "fullProfile": complete}
            if complete:
                nonempty, formulas = [], []
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            nonempty.append((cell.row, cell.column))
                        if cell.data_type == "f":
                            formulas.append(cell.coordinate)
                item["nonemptyCells"] = len(nonempty)
                item["lastDataRow"] = max((r for r, _ in nonempty), default=0)
                item["lastDataColumn"] = max((c for _, c in nonempty), default=0)
                item["formulaCells"] = formulas
            else:
                item["notice"] = "Formatted dimensions exceed the scan limit; inspect beyond the preview before mapping this source."
            sheets.append(item)
        return {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "bytes": path.stat().st_size, "sheets": sheets}
    finally:
        workbook.close()


def json_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("sources", nargs="+", type=Path)
    parser.add_argument("--max-cells", type=int, default=1_000_000)
    args = parser.parse_args()
    if args.max_cells < 1:
        parser.error("--max-cells must be positive")
    print(json.dumps({"scope": "Stored values/formulas and layout only; formulas are not evaluated and previews are not semantic validation.",
                      "sources": [profile(path, args.max_cells) for path in args.sources]},
                     indent=2, default=json_value))


if __name__ == "__main__":
    main()
