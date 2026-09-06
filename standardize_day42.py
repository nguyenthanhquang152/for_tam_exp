"""Convert MA_42D.xlsx's individual-shrimp measurements to the common summary schema."""

import argparse
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from bacteria_data import RAW_DIR, STANDARDIZED_DIR, SUMMARY_COLUMNS
from standardize_data import file_reference, write_standardized_workbook


def standardize_day42(source, output):
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve():
        raise ValueError("The standardized workbook must not overwrite the raw source.")
    workbook = load_workbook(source, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError("Expected the day-42 raw 'Sheet1' worksheet.")
        sheet = workbook["Sheet1"]
        if [sheet.cell(1, c).value for c in range(1, 5)] != ["Treatment", "Shirmp number", "Organ", "Dilution"]:
            raise ValueError("Unexpected day-42 identification headers.")
        columns = [c for c in range(5, sheet.max_column + 1) if sheet.cell(3, c).value is not None]
        if not columns or columns != list(range(5, max(columns) + 1)):
            raise ValueError("Missing or non-contiguous taxon columns.")
        if any(not isinstance(sheet.cell(r, c).value, str) for r in (1, 3) for c in columns):
            raise ValueError("Taxon and phylum headers must be text.")
        taxa = {c: str(sheet.cell(3, c).value).strip() for c in columns}
        if len(set(taxa.values())) != len(taxa):
            raise ValueError("Duplicate taxon headers after whitespace normalization.")
        phyla = {c: str(sheet.cell(1, c).value).strip() for c in columns}
        if any(value in ("", "None") for value in [*taxa.values(), *phyla.values()]):
            raise ValueError("Missing taxon or phylum header.")
        if any(c.value is not None for row in sheet.iter_rows(min_col=max(columns) + 1) for c in row):
            raise ValueError("Unexpected data beyond the declared taxon columns.")
        # Only actual merged ranges supply missing treatment and shrimp identifiers.
        merged = {}
        for area in sheet.merged_cells.ranges:
            if area.min_row >= 4:
                if area.min_col != area.max_col or area.min_col not in (1, 2):
                    raise ValueError(f"Unexpected merged measurement cells: {area}")
                for r in range(area.min_row, area.max_row + 1):
                    merged[r, area.min_col] = sheet.cell(area.min_row, area.min_col).value
        cohorts, samples, raw, issues = defaultdict(list), set(), [], []
        for r in range(4, sheet.max_row + 1):
            if not any(sheet.cell(r, c).value is not None for c in range(1, max(columns) + 1)):
                continue
            treatment = merged.get((r, 1), sheet.cell(r, 1).value)
            shrimp = merged.get((r, 2), sheet.cell(r, 2).value)
            organ, dilution = sheet.cell(r, 3).value, sheet.cell(r, 4).value
            if treatment not in ("F0", "F15", "F25", "F35", "F50") or organ not in ("G", "HP"):
                raise ValueError(f"Row {r}: unknown treatment or organ.")
            if isinstance(shrimp, bool) or not isinstance(shrimp, (int, float)) or shrimp not in (1, 2, 3, 4, 5):
                raise ValueError(f"Row {r}: invalid shrimp identifier.")
            if isinstance(dilution, bool) or not isinstance(dilution, (int, float)) or not math.isfinite(dilution) or dilution < 0:
                raise ValueError(f"Row {r}: invalid dilution.")
            tissue = "Gut" if organ == "G" else "HP"
            key = (treatment, tissue, int(shrimp))
            if key in samples:
                raise ValueError(f"Duplicate individual-shrimp record: {key}")
            samples.add(key)
            cohorts[treatment, tissue].append(r)
            for c in columns:
                cell = sheet.cell(r, c)
                value = cell.value
                if value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0):
                    raise ValueError(f"{cell.coordinate}: expected a nonnegative measurement or blank.")
                raw.append(["DOC42", treatment, int(shrimp), tissue, dilution, taxa[c], phyla[c], value, cell.coordinate])
                if value is None:
                    issues.append(["DOC42", tissue, treatment, taxa[c], int(shrimp), cell.coordinate,
                                   "Missing measurement", "Summary n_pos and freq remain unknown"])
        expected = {(t, tissue, shrimp) for t in ("F0", "F15", "F25", "F35", "F50")
                    for tissue in ("Gut", "HP") for shrimp in range(1, 6)}
        if samples != expected:
            raise ValueError(f"Incomplete or unexpected sample roster: {expected - samples}")
        summary, checks = [], []
        for (treatment, tissue), rows in sorted(cohorts.items(), key=lambda item: (int(item[0][0][1:]), item[0][1])):
            incomplete = 0
            for c in columns:
                values = [sheet.cell(r, c).value for r in rows]
                missing = sum(value is None for value in values)
                observed = sum(value is not None and value > 0 for value in values)
                incomplete += bool(missing)
                summary.append(["DOC42", tissue, taxa[c], phyla[c], treatment, len(rows),
                                None if missing else observed, None if missing else observed / len(rows),
                                missing, observed])
            checks.append([treatment, tissue, len(columns), len(rows), incomplete])
    finally:
        workbook.close()

    report = {"source": source.name, "rawPath": file_reference(source), "method": "individual_shrimp_presence", "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
              "time": "DOC42", "blankPolicy": "missing", "rawSamples": len(samples),
              "rawMeasurements": len(raw), "missingMeasurements": len(issues),
              "summaryRows": len(summary), "incompleteSummaries": sum(r[8] > 0 for r in summary)}
    contents = {
        "summary": (SUMMARY_COLUMNS, summary),
        "validation_check": (["Treatment", "Tissue", "n_taxa", "n_shrimp_unique", "n_incomplete_taxa"], checks),
        "issues_if_any": (["Time", "Tissue", "Treatment", "Taxon", "Shrimp", "SourceCell", "Issue", "Handling"], issues),
        "raw_observations": (["Time", "Treatment", "Shrimp", "Tissue", "Dilution", "Taxon", "Phylum", "Measurement", "SourceCell"], raw),
        "normalization_notes": (["Field", "Value"], [[k, v] for k, v in report.items()] + [
            ["Positive rule", "Measurement > 0; dilution does not change presence/absence."],
            ["Blank rule", "Unknown, not zero. Any blank makes the cohort/taxon percentage unknown."],
            ["Organ mapping", "G -> Gut; HP -> HP"],
            ["Identifiers", "Treatment and shrimp identifiers resolved only through actual merged ranges."],
            ["Time label", "DOC42 from the user-specified sampling source MA_42D.xlsx."],
        ]),
    }
    return write_standardized_workbook(output, contents, report)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", type=Path, default=RAW_DIR / "MA_42D.xlsx")
    parser.add_argument("--output", type=Path, default=STANDARDIZED_DIR / "MA_42D_standardized.xlsx")
    args = parser.parse_args()
    print(json.dumps(standardize_day42(args.source, args.output), indent=2))
