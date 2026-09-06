"""Small executable check: python test_plot_bacteria.py."""

from pathlib import Path
from tempfile import TemporaryDirectory

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook

from plot_bacteria import BASE, prepare_matrix, read_data, plot_figure


def main():
    source = BASE / "MA_56D_R_corrected.xlsx"
    records = read_data(source)
    assert len(records) == 150
    assert {r["Time"] for r in records} == {"DOC56"}
    assert {r["Tissue"] for r in records} == {"HP", "Gut"}
    assert {r["n_shrimp"] for r in records} == {5}

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        headers, *raw_rows = list(workbook["summary"].values)
        raw = [dict(zip(headers, row)) for row in raw_rows]
        for tissue in ("HP", "Gut"):
            taxa, treatments, matrix = prepare_matrix([r for r in records if r["Tissue"] == tissue])
            assert matrix.shape == (15, 5)
            assert treatments == ["F0", "F15", "F25", "F35", "F50"]
            for row in raw:
                if row["Tissue"] == tissue:
                    assert matrix[taxa.index(row["Taxon"]), treatments.index(row["Treatment"])] == 100 * row["n_pos"] / row["n_shrimp"]
        checks = list(workbook["validation_check"].values)
        for treatment, tissue, n_taxa, n_shrimp in checks[1:]:
            cohort = [r for r in records if r["Treatment"] == treatment and r["Tissue"] == tissue]
            assert len(cohort) == n_taxa
            assert {r["n_shrimp"] for r in cohort} == {n_shrimp}
        for _, treatment, tissue, positive, total, _ in list(workbook["issues_if_any"].values)[1:]:
            row = next(r for r in records if (r["Treatment"], r["Tissue"], r["Taxon"]) == (treatment, tissue, "Micrococcus luteus"))
            assert (row["n_pos"], row["n_shrimp"]) == (positive, total)
    finally:
        workbook.close()

    for selected in (["HP"], ["Gut"], ["HP", "Gut"]):
        figure = plot_figure(records, "DOC56", selected)
        for panel, (ax, tissue) in enumerate(zip(figure.axes, selected)):
            image = ax.images[0]
            assert image.get_clim() == (0, 100)
            labels = [label.get_text() for label in figure.axes[0].get_yticklabels()]
            for row in records:
                if row["Tissue"] == tissue:
                    assert image.get_array()[labels.index(row["Taxon"]), treatments.index(row["Treatment"])] == row["prevalence_pct"]
            assert not ax.texts, "Heatmap cells must have no text annotations."
            assert not ax.child_axes, "Do not repeat sample-size labels below columns."
            assert not any(tick.label1.get_visible() for tick in ax.xaxis.get_major_ticks())
            assert [tick.get_text() for tick in ax.get_xticklabels()] == treatments
            horizontal = [line for line in ax.lines if len(set(line.get_ydata())) == 1]
            vertical = [line for line in ax.lines if len(set(line.get_xdata())) == 1]
            assert len(horizontal) == 3, "Every taxon section needs a dashed rule."
            assert len(vertical) == 4, "Treatment columns need dashed separators."
            assert all(line.is_dashed() for line in horizontal + vertical)
            if panel == 0:
                assert all(line.get_xdata()[0] < 0 and not line.get_clip_on() for line in horizontal)
        assert sum(text.get_text() == "n = 5 shrimp per treatment and tissue."
                   for text in figure.texts) == 1
        figure.canvas.draw()
        plt.close(figure)

    # A missing record must stay missing, even when other cells are zero.
    sparse = [r for r in records if r["Tissue"] == "HP"
              and (r["Taxon"], r["Treatment"]) != ("Vibrio owensii", "F0")]
    taxa, treatments, matrix = prepare_matrix(sparse)
    assert np.isnan(matrix[taxa.index("Vibrio owensii"), treatments.index("F0")])
    assert matrix[taxa.index("Bacillus altitudinis"), treatments.index("F0")] == 0

    # Exercise malformed workbook inputs without modifying the supplied file.
    header = ["Time", "Tissue", "Taxon", "Treatment", "n_shrimp", "n_pos", "freq"]
    valid = ["DOC56", "HP", "Vibrio owensii", "F0", 5, 4, 0.8]
    cases = [[], [valid, valid]]
    for column, value in ((4, 0), (4, 5.5), (5, -1), (5, 6), (5, 1.5),
                          (6, 0.1), (6, None), (2, "")):
        bad = valid.copy()
        bad[column] = value
        cases.append([bad])
    cases.append([valid, ["DOC56", "HP", "Bacillus altitudinis", "F0", 10, 8, 0.8]])
    with TemporaryDirectory() as directory:
        path = Path(directory) / "invalid.xlsx"
        for rows in cases:
            workbook = Workbook()
            workbook.active.title = "summary"
            workbook.active.append(header)
            for row in rows:
                workbook.active.append(row)
            workbook.save(path)
            workbook.close()
            try:
                read_data(path)
            except ValueError:
                pass
            else:
                raise AssertionError(f"Invalid data accepted: {rows}")
    print("PASS: all 150 percentages, workbook checks, three plot layouts, unnumbered cells, sample-size annotation, dashed rules, missing vs zero, and invalid inputs.")


if __name__ == "__main__":
    main()
