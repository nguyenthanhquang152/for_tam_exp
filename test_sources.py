"""Check every source and independently recount all day-42 measurements."""

import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from bacteria_data import BASE, read_data
from export_web_data import DEFAULT_SOURCES, export_data
from standardize_day42 import standardize_day42


def main():
    raw_path = BASE / "MA_42D.xlsx"
    original_hash = hashlib.sha256(raw_path.read_bytes()).hexdigest()
    data = export_data([BASE / name for name in DEFAULT_SOURCES])
    assert [s["filename"] for s in data["sources"]] == list(DEFAULT_SOURCES)
    assert [s["time"] for s in data["sources"]] == ["DOC14", "DOC28", "DOC42", "DOC56"]
    assert len(data["rows"]) == 560
    assert len({r["taxon"] for r in data["rows"]}) == 16
    assert sum(r["percent"] is None for r in data["rows"]) == 25
    for source, count in zip(data["sources"], (130, 140, 140, 150)):
        assert len([r for r in data["rows"] if r["time"] == source["time"]]) == count
        assert source["sha256"] == hashlib.sha256((BASE / source["filename"]).read_bytes()).hexdigest()
    for name in (DEFAULT_SOURCES[0], DEFAULT_SOURCES[1], DEFAULT_SOURCES[3]):
        rows = read_data(BASE / name)
        for row in rows:
            expected = 100 * row["n_pos"] / row["n_shrimp"]
            actual = next(r for r in data["rows"] if (r["time"], r["tissue"], r["treatment"], r["taxon"]) ==
                          (row["Time"], row["Tissue"], row["Treatment"], row["Taxon"]))
            assert actual["percent"] == expected

    normalized = BASE / "standardized/MA_42D_standardized.xlsx"
    lookup = {(r["Treatment"], r["Tissue"], r["Taxon"]): r for r in read_data(normalized)}
    wb = load_workbook(raw_path, data_only=True)
    sheet = wb["Sheet1"]
    missing_readings = 0
    # Use the independently inspected 10-row treatment blocks and alternating organs,
    # rather than the normalizer's merged-cell lookup and cohort accumulator.
    for start in (4, 14, 24, 34, 44):
        treatment = sheet.cell(start, 1).value
        for offset, tissue in ((0, "Gut"), (1, "HP")):
            for column in range(5, 19):
                taxon = sheet.cell(3, column).value.strip()
                values = [sheet.cell(start + 2 * shrimp + offset, column).value for shrimp in range(5)]
                missing = sum(value is None for value in values)
                positives = sum(value is not None and value > 0 for value in values)
                row = lookup[treatment, tissue, taxon]
                assert row["n_shrimp"] == 5
                assert row["n_missing"] == missing
                assert row["n_pos_observed"] == positives
                assert row["n_pos"] == (None if missing else positives)
                assert row["freq"] == (None if missing else positives / 5)
                missing_readings += missing
    assert missing_readings == 28
    assert lookup["F0", "Gut", "Tenacibaculum discolor"]["n_pos_observed"] == 1
    assert lookup["F0", "HP", "Vibrio jasicida"]["n_pos_observed"] == 4
    assert lookup["F0", "HP", "Mesoflavibacter zeaxanthinifaciens"]["freq"] == 0
    wb.close()
    wb = load_workbook(normalized, data_only=True)
    assert wb["raw_observations"].max_row == 701
    assert wb["issues_if_any"].max_row == 29
    assert len(list(wb["summary"].values)) == 141
    wb.close()
    assert hashlib.sha256(raw_path.read_bytes()).hexdigest() == original_hash

    with TemporaryDirectory() as folder:
        for coordinate, value in (("E4", -1), ("E4", "ND"), ("B6", 1), ("C5", None)):
            bad = Path(folder) / "invalid.xlsx"
            wb = load_workbook(raw_path)
            wb["Sheet1"][coordinate] = value
            wb.save(bad)
            wb.close()
            try:
                standardize_day42(bad, Path(folder) / "summary.xlsx")
            except ValueError:
                pass
            else:
                raise AssertionError(f"Invalid raw source accepted: {coordinate}={value}")
    print("PASS: 560 source summaries, all 700 day-42 readings, 28 missing readings / 25 unknown percentages, source order and immutable originals.")


if __name__ == "__main__":
    main()
