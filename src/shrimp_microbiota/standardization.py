"""Standardize every raw workbook before either heatmap pipeline reads it."""

import hashlib
import io
import json
import re
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .data import BASE, DEFAULT_SOURCES, RAW_DIR, STANDARDIZED_DIR, SUMMARY_COLUMNS, read_data


def file_reference(path):
    try:
        return str(Path(path).resolve().relative_to(BASE))
    except ValueError:
        return Path(path).name


def write_standardized_workbook(output, contents, report):
    """Write the shared layout with stable archive timestamps and a content hash."""
    output = Path(output)
    headers, rows = contents["summary"]
    if tuple(headers) != SUMMARY_COLUMNS:
        raise ValueError("Standardized summary columns must follow the canonical schema.")
    rows = sorted(rows, key=lambda r: (r[0], r[1],
        [(0, int(p)) if p.isdigit() else (1, p) for p in re.split(r"(\d+)", r[4])], r[2]))
    contents = {**contents, "summary": (SUMMARY_COLUMNS, rows)}
    result = Workbook()
    result.remove(result.active)
    for name, (headers, rows) in contents.items():
        sheet = result.create_sheet(name)
        sheet.append(list(headers))
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245C50")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(62, max(12, max(len(str(c.value or "")) for c in col) + 2))
    archive = io.BytesIO()
    result.save(archive)
    result.close()
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".tmp.xlsx")
    try:
        # Rebuilds of identical data should not dirty Git just because the clock changed.
        with ZipFile(archive) as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "docProps/core.xml":
                    root = ElementTree.fromstring(content)
                    for name in ("created", "modified"):
                        element = root.find(f"{{http://purl.org/dc/terms/}}{name}")
                        if element is not None and element.text:
                            content = content.replace(element.text.encode(), b"2000-01-01T00:00:00Z")
                stable = ZipInfo(entry.filename, date_time=(2000, 1, 1, 0, 0, 0))
                stable.compress_type = ZIP_DEFLATED
                target.writestr(stable, content)
        read_data(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    report = {**report, "standardizedPath": file_reference(output),
              "standardizedSha256": hashlib.sha256(output.read_bytes()).hexdigest()}
    output.with_suffix(".report.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return report


def standardize_summary(source, output):
    """Normalize already aggregated sources without recalculating their observations."""
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve():
        raise ValueError("The standardized workbook must not overwrite the raw source.")
    records = read_data(source)
    times = {r["Time"] for r in records}
    if len(times) != 1:
        raise ValueError("Each source must contain one sampling time.")
    for row in records:
        if not isinstance(row.get("Phylum"), str) or not row["Phylum"].strip():
            raise ValueError("Each summary row must contain a phylum label.")
        row["Phylum"] = row["Phylum"].strip()
    cohorts = defaultdict(list)
    for row in records:
        cohorts[row["Treatment"], row["Tissue"]].append(row)
    checks = [[treatment, tissue, len(rows), rows[0]["n_shrimp"], sum(r["n_missing"] > 0 for r in rows)]
              for (treatment, tissue), rows in cohorts.items()]
    issues = [[r["Time"], r["Tissue"], r["Treatment"], r["Taxon"], None, None,
               f"{r['n_missing']} missing measurements; {r['n_pos_observed']} positives observed", "n_pos and freq remain unknown"]
              for r in records if r["n_missing"]]
    report = {"source": source.name, "rawPath": file_reference(source),
              "sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "time": next(iter(times)),
              "method": "summary_column_normalization", "blankPolicy": "missing",
              "summaryRows": len(records), "missingMeasurements": sum(r["n_missing"] for r in records),
              "incompleteSummaries": len(issues)}
    contents = {
        "summary": (SUMMARY_COLUMNS, [[r[column] for column in SUMMARY_COLUMNS] for r in records]),
        "validation_check": (["Treatment", "Tissue", "n_taxa", "n_shrimp_unique", "n_incomplete_taxa"], checks),
        "issues_if_any": (["Time", "Tissue", "Treatment", "Taxon", "Shrimp", "SourceCell", "Issue", "Handling"], issues),
        "normalization_notes": (["Field", "Value"], [[k, v] for k, v in report.items()] + [
            ["Summary rule", "Canonical columns and row order; labels trimmed; source counts and frequencies preserved."],
            ["Missing rule", "No invented zeros or species aliases. Missing counts default to zero only for complete source summaries."],
            ["Original audit sheets", "Preserved in the original workbook under raw/."],
        ]),
    }
    return write_standardized_workbook(output, contents, report)


def standardize_source(source, output_dir=STANDARDIZED_DIR):
    source = Path(source)
    output = Path(output_dir) / f"{source.stem}_standardized.xlsx"
    if source.name == "MA_42D.xlsx":
        # Import after the common writer exists; the raw adapter reuses this writer.
        from .day42 import standardize_day42
        return standardize_day42(source, output)
    return standardize_summary(source, output)
