"""Check every source and independently recount all day-42 measurements."""

import hashlib
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook

from shrimp_microbiota.data import BASE, RAW_DIR, SUMMARY_COLUMNS, read_data
from shrimp_microbiota.dataset import DEFAULT_SOURCES, export_data
from shrimp_microbiota.day42 import standardize_day42
from shrimp_microbiota.standardization import standardize_source


def main():
    raw_path = RAW_DIR / "MA_42D.xlsx"
    original_hash = hashlib.sha256(raw_path.read_bytes()).hexdigest()
    with patch("shrimp_microbiota.dataset.read_data", wraps=read_data) as reader:
        data = export_data([RAW_DIR / name for name in DEFAULT_SOURCES])
        assert len(reader.call_args_list) == 4
        assert all(Path(call.args[0]).parent == BASE / "standardized" for call in reader.call_args_list)
    assert [s["filename"] for s in data["sources"]] == list(DEFAULT_SOURCES)
    assert [s["time"] for s in data["sources"]] == ["DOC14", "DOC28", "DOC42", "DOC56"]
    assert len(data["rows"]) == 560
    assert len({r["taxon"] for r in data["rows"]}) == 16
    assert sum(r["percent"] is None for r in data["rows"]) == 25
    for source, count in zip(data["sources"], (130, 140, 140, 150)):
        assert len([r for r in data["rows"] if r["time"] == source["time"]]) == count
        assert source["sha256"] == hashlib.sha256((RAW_DIR / source["filename"]).read_bytes()).hexdigest()
        assert source["rawPath"] == f"raw/{source['filename']}"
        standardized = BASE / source["standardizedFilename"]
        assert source["standardizedSha256"] == hashlib.sha256(standardized.read_bytes()).hexdigest()
        workbook = load_workbook(standardized, read_only=True, data_only=True)
        assert tuple(c.value for c in workbook["summary"][1]) == SUMMARY_COLUMNS
        assert {"summary", "validation_check", "issues_if_any", "normalization_notes"} <= set(workbook.sheetnames)
        workbook.close()
        # A repeat conversion must produce identical standardized bytes.
        standardize_source(RAW_DIR / source["filename"])
        assert source["standardizedSha256"] == hashlib.sha256(standardized.read_bytes()).hexdigest()
    for name in (DEFAULT_SOURCES[0], DEFAULT_SOURCES[1], DEFAULT_SOURCES[3]):
        rows = read_data(RAW_DIR / name)
        canonical = read_data(BASE / next(s["standardizedFilename"] for s in data["sources"] if s["filename"] == name))
        assert len(rows) == len(canonical)
        key = lambda row: (row["Time"], row["Tissue"], row["Treatment"], row["Taxon"])
        by_key = {key(row): row for row in canonical}
        for row in rows:
            assert tuple(by_key[key(row)][c] for c in SUMMARY_COLUMNS) == tuple(row[c].strip() if c == "Phylum" else row[c] for c in SUMMARY_COLUMNS)
            expected = 100 * row["n_pos"] / row["n_shrimp"]
            actual = next(r for r in data["rows"] if (r["time"], r["tissue"], r["treatment"], r["taxon"]) ==
                          (row["Time"], row["Tissue"], row["Treatment"], row["Taxon"]))
            assert actual["percent"] == expected

    normalized = BASE / "standardized/MA_42D_standardized.xlsx"
    lookup = {(r["Treatment"], r["Tissue"], r["Taxon"]): r for r in read_data(normalized)}
    wb = load_workbook(raw_path, data_only=True)
    sheet = wb["Sheet1"]
    missing_readings = 0
    # Use the independently inspected 10-row treatment blocks and alternating organs,
    # rather than the normalizer's merged-cell lookup and cohort accumulator.
    for start in (4, 14, 24, 34, 44):
        treatment = sheet.cell(start, 1).value
        for offset, tissue in ((0, "Gut"), (1, "HP")):
            for column in range(5, 19):
                taxon = sheet.cell(3, column).value.strip()
                values = [sheet.cell(start + 2 * shrimp + offset, column).value for shrimp in range(5)]
                missing = sum(value is None for value in values)
                positives = sum(value is not None and value > 0 for value in values)
                row = lookup[treatment, tissue, taxon]
                assert row["n_shrimp"] == 5
                assert row["n_missing"] == missing
                assert row["n_pos_observed"] == positives
                assert row["n_pos"] == (None if missing else positives)
                assert row["freq"] == (None if missing else positives / 5)
                missing_readings += missing
    assert missing_readings == 28
    assert lookup["F0", "Gut", "Tenacibaculum discolor"]["n_pos_observed"] == 1
    assert lookup["F0", "HP", "Vibrio jasicida"]["n_pos_observed"] == 4
    assert lookup["F0", "HP", "Mesoflavibacter zeaxanthinifaciens"]["freq"] == 0
    wb.close()
    wb = load_workbook(normalized, data_only=True)
    assert wb["raw_observations"].max_row == 701
    assert wb["issues_if_any"].max_row == 29
    assert len(list(wb["summary"].values)) == 141
    wb.close()
    assert hashlib.sha256(raw_path.read_bytes()).hexdigest() == original_hash

    with TemporaryDirectory() as folder:
        for coordinate, value in (("E4", -1), ("E4", "ND"), ("B6", 1), ("C5", None)):
            bad = Path(folder) / "invalid.xlsx"
            wb = load_workbook(raw_path)
            wb["Sheet1"][coordinate] = value
            wb.save(bad)
            wb.close()
            try:
                standardize_day42(bad, Path(folder) / "summary.xlsx")
            except ValueError:
                pass
            else:
                raise AssertionError(f"Invalid raw source accepted: {coordinate}={value}")
    print("PASS: all four standardized schemas and deterministic files, 560 source summaries, 700 day-42 readings, 25 unknown percentages, ordered sources and immutable raw inputs.")


class ContractTest(unittest.TestCase):
    def test_contract(self):
        main()


if __name__ == "__main__":
    unittest.main()
