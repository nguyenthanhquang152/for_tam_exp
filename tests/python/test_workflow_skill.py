"""Behavioral checks for the skill's inspection and release-audit helpers."""

import csv
import hashlib
import importlib.util
import json
from pathlib import Path
import shutil
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook
from shrimp_microbiota.data import SUMMARY_COLUMNS, taxon_group
from shrimp_microbiota.paths import ROOT_DIR


def helper(name):
    path = ROOT_DIR / ".agents/skills/update-shrimp-heatmap/scripts" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


audit = helper("audit_release").audit
profile = helper("inspect_sources").profile


def fixture(root):
    for directory in ("raw", "standardized", "dist", "reports/python"):
        (root / directory).mkdir(parents=True)
    rows = [
        {"time": "DOC91", "tissue": "HP", "taxon": "Vibrio fixture", "treatment": "Q0",
         "total": 3, "positive": 2, "percent": 100 * 2 / 3, "missing": 0, "observedPositive": 2, "group": 0},
        {"time": "DOC91", "tissue": "Gut", "taxon": "Bacillus fixture", "treatment": "Q7",
         "total": 4, "positive": None, "percent": None, "missing": 1, "observedPositive": 1, "group": 1},
    ]
    raw = root / "raw/fixture.xlsx"
    normalized = root / "standardized/fixture_standardized.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "summary"
    sheet.append(SUMMARY_COLUMNS)
    for row in rows:
        sheet.append([row["time"], row["tissue"], row["taxon"], "Fixture phylum", row["treatment"],
                      row["total"], row["positive"], None if row["positive"] is None else row["positive"] / row["total"],
                      row["missing"], row["observedPositive"]])
        assert row["group"] == taxon_group(row["taxon"])
    workbook.save(raw)
    workbook.close()
    shutil.copy2(raw, normalized)
    raw_hash = hashlib.sha256(raw.read_bytes()).hexdigest()
    std_hash = hashlib.sha256(normalized.read_bytes()).hexdigest()
    report = {"source": raw.name, "sha256": raw_hash, "standardizedSha256": std_hash,
              "time": "DOC91", "summaryRows": 2, "incompleteSummaries": 1, "missingMeasurements": 1}
    normalized.with_suffix(".report.json").write_text(json.dumps(report))
    data = {"schemaVersion": 2, "sources": [{"filename": raw.name, "rawPath": "raw/fixture.xlsx", "sha256": raw_hash,
             "standardizedFilename": "standardized/fixture_standardized.xlsx", "standardizedSha256": std_hash,
             "time": "DOC91", "normalization": report}], "rows": rows}
    (root / "dist/data.json").write_text(json.dumps(data))
    with (root / "reports/python/plotted_data.csv").open("w", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["Time", "Tissue", "Taxon", "Treatment", "n_shrimp", "n_pos", "freq", "prevalence_pct", "n_missing", "n_pos_observed"])
        for row in rows:
            writer.writerow([row["time"], row["tissue"], row["taxon"], row["treatment"], row["total"], row["positive"],
                             None if row["positive"] is None else row["positive"] / row["total"],
                             row["percent"], row["missing"], row["observedPositive"]])
    return data


class WorkflowHelpersTest(unittest.TestCase):
    def test_audit_accepts_a_different_sparse_corpus_and_derives_counts(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            fixture(root)
            result = audit(root, ["DOC91"])
            self.assertEqual(result["sourceCount"], 1)
            self.assertEqual(result["summaryRows"], 2)
            self.assertEqual(result["expectedCombinedGridCells"], 8)
            self.assertEqual(result["unknownPercentages"], 1)
            self.assertEqual(result["unlistedCombinations"], 6)
            with self.assertRaisesRegex(ValueError, "Sampling order"):
                audit(root, ["DOC14"])

    def test_audit_rejects_raw_changes_and_stale_csv(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            fixture(root)
            raw = root / "raw/fixture.xlsx"
            original = raw.read_bytes()
            raw.write_bytes(original + b"changed")
            with self.assertRaisesRegex(ValueError, "Raw fingerprint"):
                audit(root)
            raw.write_bytes(original)
            csv_file = root / "reports/python/plotted_data.csv"
            csv_file.write_text(csv_file.read_text().replace("66.66666666666667", "0"))
            with self.assertRaisesRegex(ValueError, "CSV"):
                audit(root)

    def test_audit_rejects_path_escape_and_changed_additive_baseline(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            data = fixture(root)
            baseline = root / "baseline.json"
            baseline.write_text(json.dumps(data))
            self.assertTrue(audit(root, baseline=baseline)["additiveBaselineChecked"])
            old = json.loads(baseline.read_text())
            old["rows"][0]["positive"] = 1
            baseline.write_text(json.dumps(old))
            with self.assertRaisesRegex(ValueError, "baseline changed"):
                audit(root, baseline=baseline)
            data["sources"][0]["rawPath"] = "../outside.xlsx"
            (root / "dist/data.json").write_text(json.dumps(data))
            with self.assertRaisesRegex(ValueError, "escapes raw"):
                audit(root)

    def test_inspection_preserves_bytes_and_marks_partial_profiles(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "Source identifier"
            sheet.merge_cells("A1:A2")
            sheet["B1"] = "=1+1"
            workbook.save(path)
            workbook.close()
            original = path.read_bytes()
            result = profile(path, 100)
            self.assertEqual(result["sheets"][0]["mergedRanges"], ["A1:A2"])
            self.assertEqual(result["sheets"][0]["formulaCells"], ["B1"])
            self.assertEqual(result["sheets"][0]["preview"][0][1], "=1+1")
            self.assertFalse(profile(path, 1)["sheets"][0]["fullProfile"])
            self.assertEqual(path.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
